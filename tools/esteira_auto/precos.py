"""Preco de venda por canal, da planilha de precificacao do hub (Almir, 23/10/2026).

Fonte: G:\\Meu Drive\\TERRA CASA DECOR\\TABELAS DE PRECIFICACAO\\Precificacao - HUB.xlsx, uma aba por canal
(ML, Magalu, Shopee, Shopify), colunas SKU | DESCRICAO | PRECO VENDA. A esteira so PROPOE: o preco vai no resultado
do pedido e o Almir confere (e preenche o que faltar) na aprovacao da ficha, no site do hub.
SKU que nao esta na aba, ou com preco vazio/zero = fica de fora (em branco na tela). Arredonda em centavos.
A planilha e copiada do Drive para a pasta do lote antes de ler (ler o G: direto pode travar sem erro); copia que
demora ou falha = sem preco, a esteira segue (o Almir preenche a mao)."""
import shutil
import threading
import unicodedata
from pathlib import Path

import openpyxl

PLANILHA = Path(r"G:\Meu Drive\TERRA CASA DECOR\TABELAS DE PRECIFICAÇÃO\Precificação - HUB.xlsx")
ABAS = {"ml": "ML", "magalu": "Magalu", "shopee": "Shopee", "shopify": "Shopify"}
PRAZO_COPIA = 120


def chave(sku):
    return str(sku or "").strip().upper().replace("_", "-")


def _norm(s):
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().strip().upper()


def copiar(origem, destino, prazo=PRAZO_COPIA):
    """Copia com prazo. Devolve o caminho da copia ou None (nao achou, falhou ou demorou)."""
    destino = Path(destino)
    erro = []

    def _c():
        try:
            shutil.copyfile(origem, destino)
        except OSError as e:
            erro.append(e)

    t = threading.Thread(target=_c, daemon=True)
    t.start()
    t.join(prazo)
    if t.is_alive() or erro or not destino.is_file():
        return None
    return destino


def _valor(x):
    try:
        v = round(float(x), 2)
    except (TypeError, ValueError):
        return None
    return v if v > 0 else None


def carregar(caminho):
    """{canal: {chave_sku: preco}}. Le as colunas pelo NOME do cabecalho (SKU, PRECO VENDA), nunca pela posicao."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        nomes = {_norm(n): n for n in wb.sheetnames}
        out = {}
        for canal, aba in ABAS.items():
            nome = nomes.get(_norm(aba))
            if nome is None:
                continue
            linhas = wb[nome].iter_rows(values_only=True)
            cab = [_norm(c) for c in (next(linhas, None) or [])]
            if "SKU" not in cab or "PRECO VENDA" not in cab:
                continue
            i_sku, i_preco = cab.index("SKU"), cab.index("PRECO VENDA")
            tabela = {}
            for r in linhas:
                if not r or len(r) <= max(i_sku, i_preco) or not r[i_sku]:
                    continue
                v = _valor(r[i_preco])
                if v is not None:
                    tabela[chave(r[i_sku])] = v
            out[canal] = tabela
        return out
    finally:
        wb.close()


def do_produto(tabela, skus):
    """{canal: {sku: preco}} so com o que a planilha tem para estes SKUs (canal sem nenhum fica de fora)."""
    out = {}
    for canal, t in (tabela or {}).items():
        achados = {s: t[chave(s)] for s in skus if chave(s) in t}
        if achados:
            out[canal] = achados
    return out


def ler_para_lote(lote, origem=PLANILHA):
    """Copia a planilha para o lote e le. Qualquer problema = {} (sem preco, a esteira segue)."""
    copia = copiar(origem, Path(lote) / "precificacao-hub.xlsx")
    if copia is None:
        return {}
    try:
        return carregar(copia)
    except Exception:  # planilha aberta/corrompida/formato novo: sem preco, nunca derruba a entrega
        return {}
