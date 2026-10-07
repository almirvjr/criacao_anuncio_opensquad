"""Tradutor: planilha padrão Tiny -> estrutura de anúncios da squad ml-anuncios.

Lê o export Tiny (64 colunas, linhas pai=agrupamento + filhas=variação de cor),
agrupa as cores por "Código do pai" e emite JSON com 1 entrada por anúncio,
cada um com sua lista de variações (cor, sku, ean, preço, custo, fotos).

NÃO publica nada e NÃO toca em fotos — só estrutura os dados. O re-hosting das
fotos no bucket é feito por rehost_fotos.py (preenche `fotos_bucket`).

Uso:
    python tiny_to_esteira.py "<entrada.xls|.xlsx>" "<saida.json>" [--cubagem PATH] [--preco-provisorio VALOR]
"""
import sys
import os
import json
import re
import argparse
import openpyxl

# garante que 'esteira' seja encontrado ao rodar como script direto
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from esteira.normalizer import normalize

# header normalizado (lower, espaços colapsados) -> campo canônico
WANT = {
    "código (sku)": "sku",
    "descrição": "nome",
    "preço": "preco_venda",
    "preço de custo": "preco_custo",
    "gtin/ean": "ean",
    "descrição complementar": "descricao_complementar",
    "código do pai": "cod_pai",
    "variações": "variacoes_raw",
    "tipo do produto": "tipo",
    "marca": "marca",
    "fornecedor": "fornecedor",
    "peso bruto (kg)": "peso_bruto",
    "largura embalagem": "larg",
    "altura embalagem": "alt",
    "comprimento embalagem": "comp",
}


def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower() if s is not None else ""


def s(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # xlrd devolve numero como float: 7896835234023.0 -> "7896835234023"
    return str(v).strip() if v not in (None, "") else None


def money(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def ler_linhas(path):
    """Linhas da 1a aba. O Tiny exporta .xls (formato antigo) ou .xlsx."""
    if path.lower().endswith(".xls"):
        import xlrd
        sh = xlrd.open_workbook(path, ignore_workbook_corruption=True).sheet_by_index(0)
        return [tuple(sh.row_values(i)) for i in range(sh.nrows)]
    ws = openpyxl.load_workbook(path, data_only=True).active
    return list(ws.iter_rows(values_only=True))


def sku_chave(sku):
    return str(sku or "").strip().upper().replace("_", "-")


def carregar_cubagem(path):
    """Planilha de cubagem (regra do Almir: fonte de peso e medida). Casa por SKU."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Produtos"]
    out = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r or not r[0]:
            continue
        out[sku_chave(r[0])] = {"peso_bruto": r[3], "largura": r[4], "altura": r[5], "comprimento": r[6]}
    return out


def _vazio(x):
    try:
        return not x or float(x) == 0
    except (TypeError, ValueError):
        return True


def aplicar_cubagem(records, cub):
    """Regra do Almir: peso e medida SEMPRE da cubagem; o Tiny so fica quando o SKU
    nao esta na planilha ou o campo dela esta vazio (o Tiny vem furado)."""
    n = 0
    for rec in records:
        c = cub.get(sku_chave(rec["sku"]))
        if not c:
            continue
        mexeu = False
        if not _vazio(c["peso_bruto"]):
            rec["peso_bruto"] = c["peso_bruto"]
            mexeu = True
        dim = dict(rec.get("dim_cm") or {})
        for k in ("largura", "altura", "comprimento"):
            if not _vazio(c[k]):
                dim[k] = c[k]
                mexeu = True
        rec["dim_cm"] = dim
        if mexeu:
            rec["fonte_medida"] = "cubagem"
            n += 1
    return n


def main():
    ap = argparse.ArgumentParser(description="Planilha Tiny -> anuncios-entrada.json")
    ap.add_argument("entrada")
    ap.add_argument("saida")
    ap.add_argument("--cubagem", help="CUBAGEM DOS PRODUTOS.xlsx (fonte de peso e medida)")
    ap.add_argument("--preco-provisorio", type=float,
                    help="Sobrescreve o preco de todas as cores e marca preco_provisorio=true")
    args = ap.parse_args()
    inp, outp = args.entrada, args.saida
    rows = ler_linhas(inp)
    headers = [norm(h) for h in rows[0]]

    idx = {}
    img_cols = []
    for ci, h in enumerate(headers):
        if h in WANT:
            idx[WANT[h]] = ci
        if h.startswith("url imagem"):  # "URL imagem 1..6" + "URL imagem externa 1..10"
            img_cols.append(ci)

    records = []
    for r in rows[1:]:
        def cell(field):
            i = idx.get(field)
            return r[i] if i is not None and i < len(r) else None

        sku = s(cell("sku"))
        if not sku:
            continue
        fotos = [s(r[ci]) for ci in img_cols if ci < len(r) and s(r[ci])]
        records.append({
            "sku": sku,
            "nome": s(cell("nome")),
            "preco_venda": money(cell("preco_venda")),
            "preco_custo": money(cell("preco_custo")),
            "ean": s(cell("ean")),
            "descricao_complementar": s(cell("descricao_complementar")),
            "cod_pai": s(cell("cod_pai")),
            "variacoes_raw": s(cell("variacoes_raw")),
            "tipo": s(cell("tipo")),
            "marca": s(cell("marca")),
            "fornecedor": s(cell("fornecedor")),
            "peso_bruto": cell("peso_bruto"),
            "dim_cm": {"largura": cell("larg"), "altura": cell("alt"), "comprimento": cell("comp")},
            "fotos": fotos,
        })

    if args.cubagem:
        n = aplicar_cubagem(records, carregar_cubagem(args.cubagem))
        print(f"cubagem: peso/medida preenchidos em {n} SKU(s)", file=sys.stderr)
    if args.preco_provisorio is not None:
        for r in records:
            r["preco_venda"] = round(args.preco_provisorio, 2)
    out = normalize(records)
    for a in out["anuncios"]:
        a["preco_provisorio"] = args.preco_provisorio is not None
    out["origem"] = inp

    with open(outp, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=2)
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
